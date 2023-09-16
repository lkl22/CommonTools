# -*- coding: utf-8 -*-
# python 3.x
# Filename: Word2Excel.py
# 定义一个Word2Excel类实现word题库转Excel文件

import re
from docx import Document
from collections import OrderedDict

from openpyxl.styles import Alignment, PatternFill

from util.LogUtil import LogUtil
from util.excel.Excel2007Operator import Excel2007Operator

questionDescRule = re.compile("(\d+[．.])(.*)")
answerRule = re.compile("[\S\s]*([(（][ABCDEFGHIJK ]+[)）])[\S\s]*")

optionRule = re.compile("([ABCDEFGHIJK ]{1,2}[.、．])")

solutionRule = re.compile("(解析[：:])(.*)")

KEY_ANSWER = 'answer'


class Word2Excel:
    def __init__(self, wordFilePath='../../resources/mockExam/题库.docx', saveFn="题库.xlsx", examInfo=None):
        self.doc = Document(wordFilePath)
        self.questionDataDict = self.getWordTxt()
        self.genExcel(saveFn, examInfo)

    def getWordTxt(self):
        # 保存最终的结构化数据
        questionDataDict = OrderedDict()

        questionNo = 0
        questionObj = None
        answer = None
        txtType = None
        curOption = None
        for paragraph in self.doc.paragraphs:
            text = paragraph.text
            # 对于空白行就直接跳过
            if not text:
                continue
            # LogUtil.d("paragraph text: ", text)
            matchObj = questionDescRule.match(text)
            if matchObj:
                questionDesc = matchObj.group(2).strip()
                # LogUtil.d("match questionDescRule", matchObj.group(1)[:-1], questionDesc)
                if questionObj:
                    questionDataDict[questionNo] = questionObj

                txtType = 'questionDesc'
                questionNo += 1
                questionObj = OrderedDict()
                questionObj['question'] = questionDesc
                answer = None
                continue

            matchObj = optionRule.match(text)
            if matchObj:
                txtType = 'option'
                matchObj = optionRule.split(text)
                if not answer:
                    try:
                        answer = answerRule.match(questionObj['question']).group(1)[1:-1].strip()
                        questionObj[KEY_ANSWER] = answer
                        questionObj['question'] = re.sub("[(（][ABCDEFGHIJK ]+[)）]", "()", questionObj['question'])
                    except Exception as err:
                        LogUtil.e('Find answer error：', err, "\n", matchObj, questionObj)
                # print(matchObj)
                for txt in matchObj:
                    txt = txt.strip()
                    if txt:
                        if optionRule.match(txt):
                            curOption = txt[:1]
                            continue
                        else:
                            questionObj[curOption] = txt
                continue

            matchObj = solutionRule.match(text)
            if matchObj:
                txtType = 'solution'
                questionObj['solution'] = matchObj.group(2).strip()
                continue

            if txtType == 'questionDesc':
                questionObj['question'] += "\n" + text
            elif txtType == 'option':
                if len(text.strip()) == 0:
                    continue
                if curOption in questionObj.keys():
                    questionObj[curOption] += "\n" + text
                else:
                    questionObj[curOption] = text
            elif txtType == 'solution':
                # print("solution", questionObj['solution'], text)
                if len(questionObj['solution']) > 0:
                    questionObj['solution'] += "\n" + text
                else:
                    questionObj['solution'] = text

        if questionObj:
            questionDataDict[questionNo] = questionObj

        # for no, questionObj in questionDataDict.items():
        #     print(no, questionObj)
        return questionDataDict

    @staticmethod
    def writeSheet(st, row, col, dict, key):
        if key in dict.keys():
            Excel2007Operator.writeSheet(st, row, col, dict[key])

    def genExcel(self, saveFn, examInfo=None):
        if examInfo is None:
            examInfo = {
                "examNam": "", "scoreLine": 80, "examTime": 90, "judgmentNums": 0, "judgmentScore": 2,
                "oneChoiceNums": 0, "oneChoiceScore": 3, "multiChoiceNums": 0, "multiChoiceScore": 5
            }
        bk = Excel2007Operator.createBook()
        examInfoSheet = Excel2007Operator.addSheet(bk, "考试信息")
        judgmentSheet = Excel2007Operator.addSheet(bk, "判断题", 1)
        oneChoiceSheet = Excel2007Operator.addSheet(bk, "单选题", 2)
        multiChoiceSheet = Excel2007Operator.addSheet(bk, "多选题", 3)

        alignment = Alignment(
            horizontal='center',  # 水平对齐，可选general、left、center、right、fill、justify、centerContinuous、distributed
            vertical='center',  # 垂直对齐， 可选top、center、bottom、justify、distributed
            text_rotation=0,  # 字体旋转，0~180整数
            wrap_text=False,  # 是否自动换行
            shrink_to_fit=False,  # 是否缩小字体填充
            indent=0,  # 缩进值
        )

        examExtInfo = [{"科目名称": examInfo["examNam"]}, {"分数线": examInfo["scoreLine"]}, {"考试时长": examInfo["examTime"]},
                       {"判断题个数": examInfo["judgmentNums"]}, {"判断题分值": examInfo["judgmentScore"]},
                       {"单选题个数": examInfo["oneChoiceNums"]},
                       {"单选题分值": examInfo["oneChoiceScore"]}, {"多选题个数": examInfo["multiChoiceNums"]},
                       {"多选题分值": examInfo["multiChoiceScore"]}]
        for index in range(len(examExtInfo)):
            for key, value in examExtInfo[index].items():
                Excel2007Operator.writeSheet(examInfoSheet, 1, index + 1, key)
                Excel2007Operator.writeSheet(examInfoSheet, 2, index + 1, value)
            examInfoSheet["ABCDEFGHI"[index] + "1"].alignment = alignment
            examInfoSheet["ABCDEFGHI"[index] + "2"].alignment = alignment
            examInfoSheet.column_dimensions["ABCDEFGHI"[index]].width = 16

        for st in [judgmentSheet, oneChoiceSheet, multiChoiceSheet]:
            Excel2007Operator.writeSheet(st, 1, 1, "题干")
            st.column_dimensions["A"].width = 48
            st.column_dimensions["B"].width = 8
            for ch in "CDEFGHIJKLMNO":
                st.column_dimensions[ch].width = 16
            Excel2007Operator.writeSheet(st, 1, 2, "答案")
            Excel2007Operator.writeSheet(st, 1, 3, "备注")
            Excel2007Operator.writeSheet(st, 1, 4, "解析")
            for index, ch in enumerate("ABCDEFGHIJK"):
                Excel2007Operator.writeSheet(st, 1, 5 + index, "选项" + ch)
            for ch in "ABCDEFGHIJKLMNO":
                st[ch + "1"].alignment = alignment
                st[ch + "1"].fill = PatternFill(
                    patternType="solid",
                    # 填充类型，可选none、solid、darkGray、mediumGray、lightGray、lightDown、lightGray、lightGrid
                    fgColor="00ffff",  # 前景色，16进制rgb
                    # bgColor="00ffff",  # 背景色，16进制rgb
                    # fill_type=None,     # 填充类型
                    # start_color=None,   # 前景色，16进制rgb
                    # end_color=None      # 背景色，16进制rgb
                )

        judgmentRow = 1
        oneChoiceRow = 1
        multiChoiceRow = 1
        for no, questionObj in self.questionDataDict.items():
            if KEY_ANSWER not in questionObj:
                LogUtil.e("questionObj have not answer. No: ", no, questionObj)
                continue
            if len(questionObj[KEY_ANSWER]) > 1:
                st = multiChoiceSheet
                multiChoiceRow += 1
                row = multiChoiceRow
            else:
                if "C" not in questionObj:
                    st = judgmentSheet
                    judgmentRow += 1
                    row = judgmentRow
                else:
                    st = oneChoiceSheet
                    oneChoiceRow += 1
                    row = oneChoiceRow

            Word2Excel.writeSheet(st, row, 1, questionObj, "question")
            Word2Excel.writeSheet(st, row, 2, questionObj, "answer")
            Word2Excel.writeSheet(st, row, 4, questionObj, "solution")
            for index, ch in enumerate("ABCDEFGHIJK"):
                Word2Excel.writeSheet(st, row, 5 + index, questionObj, ch)

        bk.save(saveFn)
        bk.close()


if __name__ == '__main__':
    Word2Excel()
    # print(answerRule.match('类Test1定义如下：\npublic class Test1{\npublic float aMethod（float a，float b）{ return 0;}\n}\n将以下哪种方法插入行3 是不合法的。 (B )  '))
    # print(re.sub("[(（][ABCDEFGHIJK ]+[)）]", "()", '网络安全是在分布网络环境中对（ D ）'))
    pass
